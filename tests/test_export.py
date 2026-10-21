from io import BytesIO
from openpyxl import load_workbook

def test_export(client):
    r=client.post("/export",json={"orders":[{"order_id":"ORD-001","customer_name":"Demo","product_name":"notebook","quantity":10,"unit_price":5,"status":"READY_FOR_APPROVAL","risk_level":"LOW"},{"order_id":"ORD-002"}]})
    assert r.status_code==200 and r.content[:2]==b"PK"
    ws=load_workbook(BytesIO(r.content)).active
    assert ws.max_row==3 and ws["A2"].value=="ORD-001" and ws["D3"].value is None

def test_formula_safe(client):
    r=client.post("/export",json={"orders":[{"order_id":"=1+1","customer_name":"+SUM(A1:A2)","product_name":"@danger"}]})
    ws=load_workbook(BytesIO(r.content)).active
    for cell in (ws["A2"],ws["B2"],ws["C2"]):
        assert cell.data_type!="f" and cell.value.startswith("'")

def test_empty(client):
    assert client.post("/export",json={"orders":[]}).status_code==422

def test_invalid(client):
    assert client.post("/export",json={"orders":[{"quantity":-1}]}).status_code==422
    assert client.post("/export",json={"orders":"wrong"}).status_code==422
