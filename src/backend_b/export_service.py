from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .schemas import ExportOrder

COLUMNS = [("Order ID", "order_id"), ("Customer Name", "customer_name"),
           ("Product Name", "product_name"), ("Quantity", "quantity"),
           ("Unit Price", "unit_price"), ("Status", "status"), ("Risk Level", "risk_level")]

def safe_cell(value):
    if isinstance(value, str):
        # Excel formula injection: force untrusted strings to remain text.
        if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value

def generate_excel(orders: list[ExportOrder]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append([label for label, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
        cell.alignment = Alignment(horizontal="center")
    for order in orders:
        ws.append([safe_cell(getattr(order, field)) for _, field in COLUMNS])
    for idx, width in enumerate((20, 26, 26, 14, 16, 28, 18), 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
