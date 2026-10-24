"""Excel export for approved and analysed orders.

Adapted from Backend B's `export_service.py` (branch `backend-b`). Backend B
wrote one row per product; SmartOrder Hub uses multi-item orders, so this version
flattens each order into one row per line item while keeping the same
formula-injection protection and workbook styling.
"""

from __future__ import annotations

import math
import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PHONE_PATTERN = re.compile(r"^\+\d[\d\s()-]{7,}$")

COLUMNS = [
    ("Order ID", "order_id"),
    ("Status", "status"),
    ("Customer", "customer"),
    ("Contact", "contact"),
    ("Delivery Date", "delivery_date"),
    ("SKU", "sku"),
    ("Product", "product"),
    ("Quantity", "quantity"),
    ("Unit Price (RM)", "unit_price"),
    ("Line Total (RM)", "line_total"),
    ("Order Revenue (RM)", "order_revenue"),
    ("Order Margin", "margin"),
    ("Recommendation", "recommendation"),
    ("Approved By", "approved_by"),
    ("Created At", "created_at"),
]

CURRENCY_HEADERS = {"Unit Price (RM)", "Line Total (RM)", "Order Revenue (RM)"}


def safe_cell(value: Any) -> Any:
    """Keep untrusted strings as text so Excel never evaluates a formula."""

    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value if math.isfinite(float(value)) else ""
    if not isinstance(value, str):
        value = str(value)
    # A Malaysian phone number starts with "+60"; that is data, not a formula.
    if PHONE_PATTERN.match(value.strip()):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(
        ("\t", "\r", "\n")
    ):
        return "'" + value
    return value


def generate_order_excel(orders: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append([label for label, _ in COLUMNS])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24527A")
        cell.alignment = Alignment(horizontal="center")

    for order in orders:
        analysis = order.get("analysis") or {}
        approval = order.get("approval") or {}
        items = order.get("items") or [{}]
        for item in items:
            sheet.append(
                [
                    safe_cell(order.get("id")),
                    safe_cell(order.get("status")),
                    safe_cell(order.get("customerName")),
                    safe_cell(order.get("contact")),
                    safe_cell(order.get("requestedDeliveryDate")),
                    safe_cell(item.get("sku")),
                    safe_cell(item.get("name")),
                    item.get("quantity"),
                    item.get("unitPrice"),
                    item.get("lineTotal"),
                    analysis.get("revenue"),
                    analysis.get("margin"),
                    safe_cell(analysis.get("recommendation")),
                    safe_cell(approval.get("actor")),
                    safe_cell(order.get("createdAt")),
                ]
            )

    widths = (20, 16, 24, 22, 16, 16, 26, 10, 16, 16, 18, 14, 22, 18, 24)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            header = sheet.cell(row=1, column=cell.column).value
            if header in CURRENCY_HEADERS:
                cell.number_format = "#,##0.00"
            elif header == "Order Margin":
                cell.number_format = "0.0%"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
