"""End-to-end tests for the demo journey.

Run from `smartorder-hub/backend`:
    pytest -q
"""

from __future__ import annotations

import os
import tempfile

os.environ["SMARTORDERHUB_DATA_DIR"] = tempfile.mkdtemp(prefix="smartorder-hub-test-")
os.environ.pop("GEMINI_API_KEY", None)
os.environ.pop("DEEPSEEK_API_KEY", None)
os.environ.pop("DEEPSEEK_MODEL", None)
os.environ.pop("DEEPSEEK_BASE_URL", None)
os.environ.pop("AI_API_KEY", None)
os.environ.pop("AI_PROVIDER", None)
os.environ.pop("AI_MODEL", None)
os.environ.pop("AI_BASE_URL", None)

from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402

client = TestClient(app)


def reset() -> None:
    """Start every test from the F&B synthetic workspace.

    Switching industry resets the JSON store, so tests that switch to another
    profile cannot leak data into the next test.
    """

    response = client.post("/api/industry", json={"industry": "fnb"})
    assert response.status_code == 200, response.text


def message_from(sender: str) -> dict:
    messages = client.get("/api/messages").json()
    return next(item for item in messages if item["sender"] == sender)


def prepare(sender: str) -> tuple[dict, dict]:
    """Extract, create the order and run the rule analysis."""

    message = message_from(sender)
    response = client.post(
        f"/api/messages/{message['id']}/extract", json={"mode": "rules"}
    )
    assert response.status_code == 200, response.text
    extracted = response.json()

    response = client.post(f"/api/messages/{message['id']}/order")
    assert response.status_code == 201, response.text
    order = response.json()

    response = client.post(f"/api/orders/{order['id']}/analyze")
    assert response.status_code == 200, response.text
    return extracted, response.json()


def test_health_and_catalog() -> None:
    reset()
    health = client.get("/api/health").json()
    assert health["status"] == "ok"
    assert health["aiMode"] == "rules-only"

    catalog = client.get("/api/catalog").json()
    assert len(catalog["products"]) == 8
    assert len(catalog["suppliers"]) == 3


def test_main_order_creates_two_purchase_drafts() -> None:
    reset()
    extracted, order = prepare("Aisyah Rahman")

    parsed = extracted["parsedData"]
    assert extracted["status"] == "recognized_as_order"
    assert parsed["customerName"] == "Aisyah Rahman"
    assert len(parsed["items"]) == 2
    assert parsed["requestedDeliveryDate"]
    assert "Jalan SS2/24" in parsed["deliveryAddress"]
    assert parsed["method"] == "rules"

    analysis = order["analysis"]
    assert analysis["recommendation"] == "ACCEPT_WITH_PURCHASE"
    assert round(analysis["margin"], 3) == 0.384
    shortfalls = {item["sku"]: item["shortfall"] for item in analysis["stock"]}
    assert shortfalls["KOPI-3IN1"] == 110
    assert shortfalls["KUIH-LAPIS"] == 10

    # The risk score is a transparent weighted scorecard. No hard blocker is
    # present here, so the displayed score is the raw weighted sum.
    risk = analysis["riskScore"]
    assert 0 <= risk <= 100
    assert len(analysis["riskFactors"]) == 5
    assert abs(
        risk - sum(factor["weighted"] for factor in analysis["riskFactors"])
    ) <= 0.3

    response = client.post(
        f"/api/orders/{order['id']}/decision",
        json={"decision": "approve", "actor": "Staff", "role": "staff"},
    )
    assert response.status_code == 403

    response = client.post(
        f"/api/orders/{order['id']}/decision",
        json={
            "decision": "approve",
            "actor": "Demo Owner",
            "role": "owner",
            "note": "Stock and margin checked; proceed.",
        },
    )
    assert response.status_code == 200, response.text
    approved = response.json()["order"]
    assert approved["status"] == "approved"
    assert len(approved["purchaseDrafts"]) == 2
    assert sum(d["totalCost"] for d in approved["purchaseDrafts"]) == 302.5

    activity = client.get(f"/api/orders/{order['id']}/activity").json()
    actions = {event["action"] for event in activity}
    assert {"message_extracted", "order_created", "order_analyzed", "order_approved"} <= actions
    assert "purchase_draft_created" in actions

    csv_response = client.get("/api/orders/export.csv")
    assert csv_response.status_code == 200
    assert order["id"] in csv_response.text
    assert "Aisyah Rahman" in csv_response.text


def test_urgent_order_is_rejected_for_lead_time() -> None:
    reset()
    _, order = prepare("Daniel Tan")
    analysis = order["analysis"]
    assert analysis["recommendation"] == "REJECT"
    assert any("lead time" in reason.lower() for reason in analysis["reasons"])
    # A hard blocker floors the band: a reject can never look "Low risk".
    assert analysis["riskScore"] >= 50
    assert analysis["riskBand"] in ("High", "Critical")


def test_missing_address_requests_information() -> None:
    reset()
    _, order = prepare("Mr Lee")
    analysis = order["analysis"]
    assert analysis["recommendation"] == "REQUEST_INFO"
    assert "delivery address" in analysis["missingFields"]

    response = client.post(
        f"/api/orders/{order['id']}/decision",
        json={
            "decision": "request_info",
            "actor": "Staff",
            "role": "staff",
            "note": "Ask for the delivery address.",
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["order"]["status"] == "info_requested"
    assert "delivery address" in body["suggestedReply"]


def test_inquiry_is_not_an_order() -> None:
    reset()
    message = message_from("Farah")
    response = client.post(
        f"/api/messages/{message['id']}/extract", json={"mode": "rules"}
    )
    assert response.status_code == 200
    assert response.json()["status"] == "not_an_order"


def test_low_margin_is_rejected() -> None:
    reset()
    _, order = prepare("Kak Mimi")
    analysis = order["analysis"]
    assert analysis["recommendation"] == "REJECT"
    assert analysis["margin"] < 0.15
    assert any("margin" in reason.lower() for reason in analysis["reasons"])
    assert analysis["riskScore"] >= 50
    assert analysis["riskBand"] in ("High", "Critical")


def test_high_value_order_warns_owner() -> None:
    reset()
    _, order = prepare("Sarah Lim")
    checks = {check["code"]: check for check in order["analysis"]["checks"]}
    assert checks["order_value"]["level"] == "warn"
    assert order["analysis"]["recommendation"] == "ACCEPT_WITH_PURCHASE"


def test_override_requires_reason() -> None:
    reset()
    _, order = prepare("Kak Mimi")

    response = client.post(
        f"/api/orders/{order['id']}/decision",
        json={"decision": "approve", "actor": "Owner", "role": "owner"},
    )
    assert response.status_code == 409

    response = client.post(
        f"/api/orders/{order['id']}/decision",
        json={
            "decision": "approve",
            "actor": "Owner",
            "role": "owner",
            "override": True,
        },
    )
    assert response.status_code == 400

    response = client.post(
        f"/api/orders/{order['id']}/decision",
        json={
            "decision": "approve",
            "actor": "Owner",
            "role": "owner",
            "override": True,
            "note": "Strategic customer; accept the lower margin this once.",
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["order"]["status"] == "approved"


def test_after_sales_is_not_an_order() -> None:
    reset()
    response = client.post(
        "/api/messages",
        json={
            "source": "whatsapp",
            "sender": "Test Customer",
            "originalText": "I want a refund for the damaged kopi 3-in-1.",
        },
    )
    assert response.status_code == 201
    message = response.json()

    response = client.post(
        f"/api/messages/{message['id']}/extract", json={"mode": "rules"}
    )
    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "not_an_order"
    assert "After-sales" in body["parsedData"]["analysisSummary"]


def test_excel_export_rows_and_formula_safety() -> None:
    reset()
    _, order = prepare("Aisyah Rahman")

    response = client.get("/api/orders/export.xlsx")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"

    from io import BytesIO

    from openpyxl import load_workbook

    worksheet = load_workbook(BytesIO(response.content)).active
    headers = [cell.value for cell in worksheet[1]]
    assert "Order ID" in headers
    assert "Order Margin" in headers
    # Header + one row per line item (Aisyah has two).
    assert worksheet.max_row == 3
    assert worksheet["A2"].value == order["id"]
    assert worksheet["D2"].value == "+60 12-345 6789"

    from app.excel_export import generate_order_excel

    content = generate_order_excel(
        [
            {
                "id": "ORD-X",
                "status": "approved",
                "customerName": "=1+1",
                "contact": "+SUM(A1:A2)",
                "requestedDeliveryDate": "2026-10-25",
                "items": [
                    {
                        "sku": "KOPI",
                        "name": "@danger",
                        "quantity": 1,
                        "unitPrice": 1,
                        "lineTotal": 1,
                    }
                ],
                "analysis": {
                    "revenue": 1,
                    "margin": 0.5,
                    "recommendation": "ACCEPT",
                },
                "approval": {"actor": "=cmd"},
                "createdAt": "2026-10-21",
            }
        ]
    )
    worksheet = load_workbook(BytesIO(content)).active
    for cell in ("C2", "D2", "G2", "N2"):
        assert worksheet[cell].value.startswith("'")
        assert worksheet[cell].data_type == "s"


def test_intake_process_builds_clean_review_table() -> None:
    reset()
    response = client.post("/api/intake/process")
    assert response.status_code == 200
    result = response.json()
    assert result["processed"] == 34
    assert result["ordersCreated"] == 15
    assert result["byType"] == {
        "order": 15,
        "inquiry": 4,
        "ad": 4,
        "after_sales": 4,
        "other": 4,
        "uncertain": 3,
    }

    summary = client.get("/api/intake/summary").json()
    assert summary["totalMessages"] == 34
    assert summary["unprocessed"] == 0
    assert summary["orders"] == 15
    assert summary["awaitingConfirmation"] == 15

    orders = client.get("/api/orders").json()
    names = {order["customerName"] for order in orders}
    assert "Aisyah Rahman" in names
    assert "Daniel Tan" in names
    assert "Farah" not in names
    assert "Unknown" not in names
    assert all(order["analysis"] for order in orders)
    assert all(order["source"] in ("whatsapp", "email") for order in orders)

    messages = client.get("/api/messages").json()
    ads = [
        message
        for message in messages
        if (message.get("parsedData") or {}).get("messageType") == "ad"
    ]
    assert len(ads) == 4
    assert all(message["status"] == "not_an_order" for message in ads)


def test_demo_reset_clears_orders() -> None:
    reset()
    response = client.post("/api/demo/reset")
    assert response.status_code == 200, response.text
    assert client.get("/api/orders").json() == []
    assert len(client.get("/api/messages").json()) == 34


def test_industry_profiles_are_complete() -> None:
    reset()
    industries = client.get("/api/industries").json()["industries"]
    assert {profile["id"] for profile in industries} == {
        "fnb",
        "retail",
        "wholesale",
        "services",
        "manufacturing",
        "logistics",
    }
    for profile in industries:
        assert profile["name"]
        assert profile["fields"]
        assert profile["requiredFields"] is not None
        assert set(profile["riskWeights"]) == {
            "stock",
            "delivery",
            "margin",
            "data",
            "anomaly",
        }
        assert abs(sum(profile["riskWeights"].values()) - 1.0) < 1e-9


def test_industry_switch_changes_dataset_and_order_fields() -> None:
    reset()
    response = client.post("/api/industry", json={"industry": "wholesale"})
    assert response.status_code == 200, response.text
    workspace = response.json()["workspace"]
    assert workspace["industry"]["id"] == "wholesale"
    assert workspace["industry"]["labels"]["summary"] == "Bulk items"
    assert workspace["stats"]["totalMessages"] == 10
    assert workspace["stats"]["orders"] == 0

    catalog = client.get("/api/catalog").json()
    assert catalog["industry"]["id"] == "wholesale"
    assert len(catalog["products"]) == 4
    assert len(catalog["suppliers"]) == 2

    result = client.post("/api/intake/process").json()
    assert result["processed"] == 10
    assert result["ordersCreated"] == 6

    orders = client.get("/api/orders").json()
    assert all(order["industry"] == "wholesale" for order in orders)
    assert all(order["analysis"] for order in orders)
    assert any(order["industryFields"].get("poNumber") for order in orders)
    assert all(
        order["analysis"]["riskScore"] >= 50
        for order in orders
        if order["analysis"]["recommendation"] in ("REJECT", "NEEDS_REVIEW")
    )

    # Leave the shared test store back on F&B for any later test.
    reset()


def test_ai_status_and_connection_guard() -> None:
    reset()
    status = client.get("/api/ai/status").json()
    assert status["configured"] is False
    assert status["provider"] == "rules"
    assert status["keyEnv"] == "DEEPSEEK_API_KEY"
    assert status["intakeMode"] in ("rules", "auto", "ai")

    response = client.post("/api/ai/test")
    assert response.status_code == 400
    assert "No AI key configured" in response.json()["detail"]


def test_channel_collection_sync_and_webhook() -> None:
    reset()
    channels = client.get("/api/channels").json()
    assert {channel["id"] for channel in channels} == {
        "CH-WHATSAPP",
        "CH-FACEBOOK",
        "CH-INSTAGRAM",
        "CH-TELEGRAM",
        "CH-EMAIL",
        "CH-TIKTOK",
    }
    facebook = next(channel for channel in channels if channel["id"] == "CH-FACEBOOK")
    assert facebook["status"] == "needs_attention"

    response = client.post("/api/channels/CH-FACEBOOK/connect")
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "connected"

    before = len(client.get("/api/messages").json())
    response = client.post("/api/channels/CH-FACEBOOK/sync")
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["channel"]["importedCount"] == 1
    assert len(result["messages"]) == 1
    message = result["messages"][0]
    assert message["source"] == "facebook"
    assert message["channelId"] == "CH-FACEBOOK"
    assert message["status"] == "not_analyzed"
    assert len(client.get("/api/messages").json()) == before + 1

    response = client.post("/api/channels/sync-all")
    assert response.status_code == 200, response.text
    synced = response.json()
    assert synced["syncedChannels"] == 5
    assert len(synced["messages"]) == 5

    webhook_payload = {
        "sender": "Instagram order",
        "contact": "@demo.customer",
        "originalText": "Can I order 10 packs of nasi lemak for tomorrow?",
        "externalId": "demo-webhook-001",
    }
    response = client.post(
        "/api/channels/CH-INSTAGRAM/webhook", json=webhook_payload
    )
    assert response.status_code == 201, response.text
    webhook_message = response.json()["message"]
    assert webhook_message["source"] == "instagram"
    assert webhook_message["channelAccount"] == "@kopikita.my"

    after_first_webhook = len(client.get("/api/messages").json())
    response = client.post(
        "/api/channels/CH-INSTAGRAM/webhook",
        json={**webhook_payload, "externalId": "demo-webhook-001"},
    )
    assert response.status_code == 201
    assert response.json()["message"]["id"] == webhook_message["id"]
    assert len(client.get("/api/messages").json()) == after_first_webhook

    response = client.post(
        "/api/channels/CH-TIKTOK/webhook",
        json={**webhook_payload, "externalId": "demo-webhook-002"},
    )
    assert response.status_code == 409



def test_ai_provider_detection() -> None:
    from app.config import detect_provider

    assert detect_provider("gemini", deepseek_key="ds-key") == "deepseek"
    assert detect_provider("gemini", gemini_key="gem-key") == "gemini"
    assert detect_provider("auto", deepseek_key="ds-key") == "deepseek"
    assert detect_provider("auto", gemini_key="gem-key") == "gemini"
    assert (
        detect_provider(
            "openai",
            base_url="https://api.deepseek.com",
            model="deepseek-flash",
        )
        == "deepseek"
    )
    assert detect_provider("auto") == "deepseek"


def test_risk_insights_rule_brief() -> None:
    reset()
    client.post("/api/intake/process")
    response = client.get("/api/risk/insights?language=en")
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["generatedBy"] == "rules"
    assert data["aiConfigured"] is False
    assert len(data["topRisks"]) == 2
    assert data["topRisks"][0]["share"] >= data["topRisks"][1]["share"]
    assert abs(sum(item["share"] for item in data["allRisks"]) - 1.0) < 0.01
    assert data["summary"]
    assert "%%" not in data["summary"]
    assert len(data["suggestions"]) == 2

    zh = client.get("/api/risk/insights?language=zh").json()
    assert zh["language"] == "zh"
    assert any("一" <= char <= "鿿" for char in zh["summary"])
    assert "%%" not in zh["summary"]


def test_risk_insights_uses_ai_when_configured(monkeypatch) -> None:
    from app import ai_client, config

    reset()
    client.post("/api/intake/process")
    monkeypatch.setattr(config, "AI_API_KEY", "test-key")
    monkeypatch.setattr(config, "AI_PROVIDER", "deepseek")
    monkeypatch.setattr(config, "AI_MODEL", "deepseek-flash")

    def fake_generate(prompt: str) -> str:
        assert "two largest risk types" in prompt
        return '{"summary":"Focus on stock first.","suggestions":["Review safety stock.","Confirm lead times."]}'

    monkeypatch.setattr(ai_client, "generate_json", fake_generate)
    data = client.get("/api/risk/insights?language=en").json()
    assert data["generatedBy"] == "ai"
    assert data["providerLabel"] == "DeepSeek Flash"
    assert data["summary"] == "Focus on stock first."
    assert data["suggestions"] == ["Review safety stock.", "Confirm lead times."]


def test_deepseek_payload_disables_thinking(monkeypatch) -> None:
    from app import ai_client, config

    captured: dict = {}

    class FakeResponse:
        def raise_for_status(self) -> None:
            return None

        def json(self) -> dict:
            return {"choices": [{"message": {"content": '{"ok": true}'}}]}

    class FakeClient:
        def __init__(self, timeout=None) -> None:
            captured["timeout"] = timeout

        def __enter__(self):
            return self

        def __exit__(self, *args) -> bool:
            return False

        def post(self, url, headers=None, json=None):
            captured.update(url=url, headers=headers, json=json)
            return FakeResponse()

    monkeypatch.setattr(config, "AI_PROVIDER", "deepseek")
    monkeypatch.setattr(config, "AI_API_KEY", "test-key")
    monkeypatch.setattr(config, "AI_MODEL", "deepseek-flash")
    monkeypatch.setattr(config, "AI_BASE_URL", "https://api.deepseek.com")
    monkeypatch.setattr(ai_client.httpx, "Client", FakeClient)

    assert ai_client.generate_json("extract this") == '{"ok": true}'
    assert captured["url"] == "https://api.deepseek.com/chat/completions"
    assert captured["json"]["model"] == "deepseek-flash"
    assert captured["json"]["thinking"] == {"type": "disabled"}
    assert captured["json"]["response_format"] == {"type": "json_object"}
